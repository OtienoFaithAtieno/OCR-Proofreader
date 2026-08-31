from docx import Document as WordDocument

from app.cleaner.cleaner import clean_document
from app.cleaner.rules.validations import validate_output_document
from app.core.pipeline import process_docx
from app.model.document import Document, DocumentBlock, Page


def test_output_validation_rules_normalize_document_defaults():
    page = Page(number=1, width=612, height=792)
    page.add_block(
        DocumentBlock(
            text="One  two\tthree—four\fFive . . six",
            style={"font_name": "Arial", "font_size": "9"},
        )
    )
    document = Document(pages=[page])

    clean_document(document)

    block = page.blocks[0]
    assert block.text == "One two three--four\n\nFive .. six"
    assert block.style == {
        "font_name": "Times New Roman",
        "font_size": "12.0",
        "text_color": "000000",
        "background_color": "FFFFFF",
        "style": "Normal",
    }
    assert document.properties["background_color"] == "FFFFFF"
    assert validate_output_document(document)


def test_process_docx_applies_cleaning_pipeline(tmp_path):
    docx_path = tmp_path / "raw.docx"
    word_document = WordDocument()
    word_document.add_paragraph("One  two\t. . three")
    word_document.add_paragraph("Final")
    word_document.save(docx_path)

    result = process_docx(docx_path)
    blocks = result["document"].pages[0].blocks

    assert [block.text for block in blocks] == ["One two .. three", "Final"]
    assert blocks[0].style["font_name"] == "Times New Roman"
    assert blocks[0].style["font_size"] == "12.0"
    assert blocks[0].style["text_color"] == "000000"
    assert blocks[0].style["background_color"] == "FFFFFF"