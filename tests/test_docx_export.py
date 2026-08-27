from pathlib import Path
import fitz
from docx import Document as WordDocument

from app.core.pipeline import process_pdf
from app.exporter.exporter import export_document_to_docx


def test_export_document_to_docx_contains_editable_text_only(tmp_path: Path) -> None:
    pdf_path = tmp_path / "source.pdf"
    docx_path = tmp_path / "output.docx"

    pdf = fitz.open()
    for page_number in range(2):
        page = pdf.new_page(width=600, height=800)
        page.insert_text((72, 72), f"Page {page_number + 1}")
    pdf.save(pdf_path)
    pdf.close()

    result = process_pdf(pdf_path)
    result["document"].pages[0].blocks[0].text = "Edited page 1"
    export_document_to_docx(result["document"], docx_path)

    assert docx_path.exists()
    assert len(WordDocument(docx_path).sections) == 2
    word_document = WordDocument(docx_path)
    assert word_document.core_properties.title == ""
    assert word_document.core_properties.author == ""
    assert [
        paragraph.text
        for paragraph in word_document.paragraphs
        if paragraph.text
    ] == ["Edited page 1", "Page 2"]

    assert all(
        relationship.reltype.endswith("/image") is False
        for relationship in word_document.part.rels.values()
    )
