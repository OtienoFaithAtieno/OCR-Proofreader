from pathlib import Path

from docx import Document as WordDocument

from app.core.pipeline import process_docx


def test_process_docx_loads_unformatted_text_without_metadata(tmp_path: Path) -> None:
    docx_path = tmp_path / "source.docx"
    source = WordDocument()
    source.add_paragraph("Unformatted paragraph")
    source.add_paragraph("Second paragraph")
    source.core_properties.author = "Should not be imported"
    source.core_properties.title = "Should not be imported"
    source.save(docx_path)

    result = process_docx(docx_path)

    assert result["text"] == "Unformatted paragraph\nSecond paragraph"
    assert result["document"].metadata.author == ""
    assert result["document"].metadata.title == ""
    assert result["document"].pages[0].blocks[1].text == "Second paragraph"
