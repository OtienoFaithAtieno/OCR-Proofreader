"""Simple document processing pipeline for PDFs."""

from pathlib import Path

import fitz
from docx import Document as WordDocument

from app.cleaner.cleaner import clean_document
from app.model.document import (
    BlockType,
    Document,
    DocumentBlock,
    DocumentMetadata,
    Page as DocumentPage,
)
from app.pdf.loader import PDFLoader
from app.ocr.ocr_engine import OCRUnavailableError, extract_page_text


def process_pdf(document_path: str | Path) -> dict:
    """Open a PDF, extract text, and return metadata for the UI."""
    path = Path(document_path)
    loader = PDFLoader()
    loader.open(path)

    try:
        metadata = loader.metadata
        pages: list[DocumentPage] = []
        text_chunks: list[str] = []

        for page_number in range(loader.page_count):
            page = loader.get_page(page_number)
            text = page.get_text("text").strip()
            if not text:
                try:
                    text = extract_page_text(page)
                except OCRUnavailableError as exc:
                    text = f"[Scanned page: {exc}]"
            text_chunks.append(text)

            blocks = []
            for line in text.splitlines():
                stripped = line.strip()
                if stripped:
                    blocks.append(
                        DocumentBlock(
                            type=BlockType.PARAGRAPH,
                            text=stripped,
                            confidence=0.95,
                            metadata={"source": "ocr" if not page.get_text("text").strip() else "pdf"},
                        )
                    )

            doc_page = DocumentPage(
                number=page_number + 1,
                width=float(page.rect.width),
                height=float(page.rect.height),
                metadata={"source": "pdf"},
            )
            doc_page.blocks = blocks
            pages.append(doc_page)

        document_model = Document(
            filename=path.name,
            source_path=str(path),
            metadata=DocumentMetadata(
                title="",
                author="",
                language="en",
            ),
            pages=pages,
        )
        document_model = clean_document(document_model)
        text = "\n\n".join(chunk for chunk in text_chunks if chunk)

        return {
            "document": document_model,
            "text": text,
            "metadata": {
                "page_count": document_model.page_count,
                "path": str(path),
            },
        }
    finally:
        loader.close()


def run_pipeline(document_path: str | Path) -> dict:
    """Compatibility wrapper for the old pipeline name."""
    return process_pdf(document_path)


def process_docx(document_path: str | Path) -> dict:
    """Load an unformatted DOCX into the application document model."""
    path = Path(document_path)
    word_document = WordDocument(path)
    page = DocumentPage(number=1, width=612.0, height=792.0, metadata={"source": "docx"})

    for paragraph in word_document.paragraphs:
        text = paragraph.text.strip()
        if text:
            page.add_block(
                DocumentBlock(
                    type=BlockType.PARAGRAPH,
                    text=text,
                    confidence=1.0,
                    metadata={"source": "docx"},
                )
            )

    document_model = Document(
        filename=path.name,
        source_path=str(path),
        metadata=DocumentMetadata(title="", author="", language="en"),
        pages=[page],
    )
    return {
        "document": document_model,
        "text": document_model.full_text,
        "metadata": {"page_count": 1, "path": str(path)},
    }
