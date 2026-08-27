"""Document export helpers."""

from __future__ import annotations

from pathlib import Path

from docx import Document as WordDocument
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def export_document_to_docx(document_model, output_path: str | Path) -> Path:
    """Export the document model as a fully editable text-only DOCX."""
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    word_document = WordDocument()
    word_document.core_properties.title = ""
    word_document.core_properties.author = ""
    word_document.core_properties.subject = ""
    word_document.core_properties.keywords = ""

    for page_index, model_page in enumerate(document_model.pages):
        if page_index:
            word_document.add_section(WD_SECTION.NEW_PAGE)

        section = word_document.sections[-1]
        section.page_width = Inches(model_page.width / 72)
        section.page_height = Inches(model_page.height / 72)
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

        for block in model_page.blocks:
            paragraph = word_document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(block.text)
            run.font.name = "Calibri"
            run.font.size = Pt(11)

    word_document.save(output)
    return output
